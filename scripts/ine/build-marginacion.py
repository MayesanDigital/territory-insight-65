"""
Construye el índice de marginación por sección electoral de Zacatecas y extrae
la pobreza municipal oficial de CONEVAL.

Por qué un índice propio: la medición oficial de pobreza (CONEVAL) solo llega a
municipio, y los índices de CONAPO se publican por AGEB y localidad, que no
coinciden con las secciones del INE. El ECEG 2020 (Censo 2020 asignado a
sección por INE e INEGI) sí trae las carencias que usa CONAPO, así que el índice
se calcula con el mismo método:

- Indicadores de carencia en porcentaje (más alto = peor).
- Síntesis con la distancia DP2 de Pena Trapero, que pondera cada indicador por
  la información que no aportan los anteriores (1 - R²) y así no cuenta dos
  veces carencias correlacionadas.
- Índice normalizado 0–100 y cinco grados por estratificación óptima de
  Dalenius–Hodges.

Entradas:
  data/ine/ECEG_32_Zacatecas.xlsx
  data/coneval/Concentrado_indicadores_de_pobreza_2020.xlsx
    (https://www.coneval.org.mx/Medicion/Documents/Pobreza_municipal/2020/Concentrado_indicadores_de_pobreza_2020.zip)

Salidas:
  data/ine/marginacion-secciones.json
  data/coneval/pobreza-municipal-32.json

Uso:  python scripts/ine/build-marginacion.py
"""

import json
import unicodedata
from pathlib import Path

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
ECEG = ROOT / "data/ine/ECEG_32_Zacatecas.xlsx"
CONEVAL = ROOT / "data/coneval/Concentrado_indicadores_de_pobreza_2020.xlsx"
OUT_SECCIONES = ROOT / "data/ine/marginacion-secciones.json"
OUT_MUNICIPIOS = ROOT / "data/coneval/pobreza-municipal-32.json"

MIN_VIVIENDAS = 20
GRADOS = ["muy_bajo", "bajo", "medio", "alto", "muy_alto"]


def ratio(num: float, den: float) -> float:
    return 100.0 * num / den if den > 0 else 0.0


# Carencias por sección; todas crecen con la marginación.
def indicadores(g):
    viv = g("VIVPARH_CV") or g("TVIVPARHAB")
    return {
        "no_asiste_6_14": ratio(g("P6A11_NOA") + g("P12A14NOA"), g("P_6A11") + g("P_12A14")),
        "sin_basica_15": ratio(
            g("P15YM_SE") + g("P15PRI_IN") + g("P15PRI_CO") + g("P15SEC_IN"), g("P_15YMAS")
        ),
        "sin_salud": ratio(g("PSINDER"), g("POBTOT")),
        "sin_drenaje": ratio(g("VPH_NODREN"), viv),
        "sin_electricidad": ratio(g("VPH_S_ELEC"), viv),
        "sin_agua": ratio(g("VPH_AGUAFV"), viv),
        "piso_tierra": ratio(g("VPH_PISOTI"), viv),
        "ocupantes_cuarto": g("PRO_OCUP_C"),
        "sin_refrigerador": ratio(viv - g("VPH_REFRI"), viv),
        "sin_internet": ratio(viv - g("VPH_INTER"), viv),
    }


def r2(y: np.ndarray, X: np.ndarray) -> float:
    """Coeficiente de determinación de y sobre las columnas de X (con constante)."""
    A = np.column_stack([np.ones(len(y)), X])
    coef, *_ = np.linalg.lstsq(A, y, rcond=None)
    resid = y - A @ coef
    tot = ((y - y.mean()) ** 2).sum()
    return 0.0 if tot == 0 else max(0.0, min(1.0, 1 - (resid**2).sum() / tot))


def dp2(M: np.ndarray) -> np.ndarray:
    """Distancia DP2 respecto de la mejor situación observada en cada indicador."""
    sigma = M.std(axis=0)
    sigma[sigma == 0] = 1
    D = np.abs(M - M.min(axis=0)) / sigma
    indice = D.sum(axis=1)  # Distancia de Fréchet como punto de partida.
    orden_prev = None
    for _ in range(50):
        corr = [abs(np.corrcoef(M[:, i], indice)[0, 1]) if M[:, i].std() > 0 else 0 for i in range(M.shape[1])]
        orden = list(np.argsort(corr)[::-1])
        nuevo = D[:, orden[0]].copy()
        for k in range(1, len(orden)):
            nuevo += D[:, orden[k]] * (1 - r2(M[:, orden[k]], M[:, orden[:k]]))
        if orden == orden_prev and np.allclose(nuevo, indice):
            break
        indice, orden_prev = nuevo, orden
    return indice


def dalenius_hodges(x: np.ndarray, estratos: int = 5, clases: int = 100) -> np.ndarray:
    """Cortes de estratificación óptima por raíz acumulada de frecuencias."""
    freq, edges = np.histogram(x, bins=clases)
    acum = np.cumsum(np.sqrt(freq))
    paso = acum[-1] / estratos
    cortes = []
    for s in range(1, estratos):
        i = int(np.argmin(np.abs(acum - paso * s)))
        cortes.append(edges[i + 1])
    return np.array(cortes)


def secciones():
    wb = openpyxl.load_workbook(ECEG, read_only=True)
    filas = list(wb.worksheets[0].iter_rows(values_only=True))
    h = {k: i for i, k in enumerate(filas[0])}

    registros = []
    for f in filas[1:]:
        g = lambda k: float(f[h[k]] or 0)  # noqa: E731
        # Con muy pocas viviendas un solo hogar mueve los porcentajes decenas de
        # puntos y la sección encabezaría la lista por azar (había una de 4
        # habitantes en segundo lugar). Quedan fuera y la app las muestra sin dato.
        if (g("VIVPARH_CV") or g("TVIVPARHAB")) < MIN_VIVIENDAS:
            continue
        registros.append(
            {
                "section_code": str(int(f[h["SECCION"]])).zfill(4),
                "population": int(g("POBTOT")),
                "indicators": indicadores(g),
            }
        )

    claves = list(registros[0]["indicators"].keys())
    M = np.array([[r["indicators"][k] for k in claves] for r in registros])
    valor = dp2(M)
    cortes = dalenius_hodges(valor)
    norm = 100 * (valor - valor.min()) / (valor.max() - valor.min())

    orden = np.argsort(-valor)
    rank = np.empty(len(valor), dtype=int)
    rank[orden] = np.arange(1, len(valor) + 1)

    for i, r in enumerate(registros):
        r["dp2"] = round(float(valor[i]), 6)
        r["index_value"] = round(float(norm[i]), 4)
        r["grade"] = GRADOS[int(np.searchsorted(cortes, valor[i], side="right"))]
        r["state_rank"] = int(rank[i])
        r["indicators"] = {k: round(v, 2) for k, v in r["indicators"].items()}

    OUT_SECCIONES.write_text(json.dumps(registros, ensure_ascii=False), encoding="utf-8")
    conteo = {g: sum(1 for r in registros if r["grade"] == g) for g in GRADOS}
    print(f"Secciones con índice: {len(registros)}  grados: {conteo}")
    mas = sorted(registros, key=lambda r: r["state_rank"])[:5]
    print("Más marginadas:", [(r["section_code"], r["index_value"], r["grade"]) for r in mas])


def normaliza(s: str) -> str:
    return unicodedata.normalize("NFD", s).encode("ascii", "ignore").decode().upper().strip()


def municipios():
    wb = openpyxl.load_workbook(CONEVAL, read_only=True)
    filas = list(wb["Concentrado municipal"].iter_rows(values_only=True))
    # Encabezados en dos renglones: 5 (grupo) y 6 (año). Se leen por posición.
    out = []
    for f in filas[8:]:
        if f[1] != "32" or not f[3]:
            continue
        out.append(
            {
                "municipio": f[4],
                "municipio_key": normaliza(f[4]),
                "municipio_code": f[3][2:],
                "population": int(f[7]),
                "poverty_pct": round(float(f[10]), 2),
                "poverty_people": int(f[13]),
                "extreme_poverty_pct": round(float(f[19]), 2),
                "extreme_poverty_people": int(f[22]),
                "moderate_poverty_pct": round(float(f[28]), 2),
            }
        )
    OUT_MUNICIPIOS.write_text(json.dumps(out, ensure_ascii=False), encoding="utf-8")
    peor = max(out, key=lambda m: m["poverty_pct"])
    print(f"Municipios CONEVAL: {len(out)}  mayor pobreza: {peor['municipio']} {peor['poverty_pct']}%")


if __name__ == "__main__":
    secciones()
    municipios()
